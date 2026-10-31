"""身材数据 Excel 导入导出底层模块（仅身材模块使用：基础记录 + 维度记录）。

模板约定（与 docs/body-import-format.md 一致，第一行为表头）：
- 基础记录：日期 | 身高(cm) | 体重(kg) | 体脂率(%)
- 维度记录：日期 | 颈围(cm) | 肩宽(cm) | 臂围(cm) | 胸围(cm) | 腰围(cm) | 臀围(cm) | 大腿围(cm) | 小腿围(cm)

导入语义与页面新增一致：留空列沿用该字段最近一次非空值；同日重复导入 = 覆盖该日该组数据；
行级校验失败只记录该行错误，不阻断其他行。goal / target 两组不参与导入导出。
"""
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.datetime import from_excel

from app.core.constants import BODY_FIELDS, BODY_GROUPS

# 导入/导出的数据分组（仅基础记录 + 维度记录）
IMPORT_GROUPS = ("basic", "girth")
SHEET_NAMES = {"basic": "基础记录", "girth": "维度记录"}
DATE_HEADER = "日期"
DATE_FORMATS = ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日")

# 合理物理范围（身高 cm / 体重 kg / 体脂 % / 围度 cm），越界即该列报错
RANGES: dict[str, tuple[float, float]] = {
    "height": (30, 250),
    "weight": (2, 300),
    "body_fat": (1, 60),
    "neck": (10, 100),
    "shoulder_width": (15, 150),
    "arm": (5, 100),
    "chest": (20, 250),
    "waist": (20, 250),
    "hip": (20, 250),
    "thigh": (10, 150),
    "calf": (5, 100),
}


def header_of(group: str) -> list[str]:
    """组 → 模板表头（日期 + 中文标签(单位)）。"""
    return [DATE_HEADER] + [f"{BODY_FIELDS[k][0]}({BODY_FIELDS[k][1]})" for k in BODY_GROUPS[group]]


def _column_label(key: str) -> str:
    """指标 key（或 record_date）→ 表头文本，用于错误明细。"""
    if key == "record_date":
        return DATE_HEADER
    label, unit, _col, _grp = BODY_FIELDS[key]
    return f"{label}({unit})"


# ---------- 模板 ----------

def build_template() -> BytesIO:
    """生成可下载导入模板：基础记录 / 维度记录 数据 Sheet（仅表头）+ 填写说明。"""
    buf = BytesIO()
    wb = Workbook()
    ws_basic = wb.active
    ws_basic.title = SHEET_NAMES["basic"]
    ws_basic.append(header_of("basic"))
    ws_girth = wb.create_sheet(SHEET_NAMES["girth"])
    ws_girth.append(header_of("girth"))
    ws_help = wb.create_sheet("填写说明")
    help_lines = [
        "1. 每行代表一个日期的记录；第一行表头请勿修改，请勿新增列。",
        "2. 日期格式：YYYY-MM-DD（例如 2026-10-31）。",
        "3. 只填写你有的列；留空的列会自动沿用该指标最近一次记录的值。",
        "4. 每行至少填写一个指标，否则该行导入失败。",
        "5. 同一日期重复导入 = 覆盖该日该组数据（以最后一次为准）。",
        "6. 数值范围校验：身高 30-250cm，体重 2-300kg，体脂率 1-60%，围度 5-250cm。",
        "7. 示例：基础记录写入「2026-10-31、170.5、65.2、（体脂留空）」，表示当天记录了身高体重，体脂沿用上次。",
    ]
    for line in help_lines:
        ws_help.append([line])

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="E8EEF4")
    for ws in (ws_basic, ws_girth):
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
        ws.freeze_panes = "A2"
        for i, col in enumerate(ws.columns, start=1):
            max_len = max((len(str(c.value)) for c in col if c.value is not None), default=6)
            ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 12), 16)

    wb.save(buf)
    return buf


# ---------- 解析 ----------

def parse_workbook(content: bytes) -> dict[str, list[dict]]:
    """解析导入文件 → {group: [{sheet, row, cells}]}，cells 为表头 → 原始单元格值。

    未知列忽略；「日期」列缺失或未识别到任何指标列时抛 ValueError（结构性错误）。
    空行跳过。日期/数值的格式校验在 validate_rows 中逐行进行。
    """
    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        result: dict[str, list[dict]] = {}
        for group in IMPORT_GROUPS:
            sheet_name = SHEET_NAMES[group]
            if sheet_name not in wb.sheetnames:
                raise ValueError(f"导入文件缺少工作表「{sheet_name}」")
            ws = wb[sheet_name]
            rows_iter = ws.iter_rows(values_only=True)
            try:
                header = next(rows_iter)
            except StopIteration:
                raise ValueError(f"工作表「{sheet_name}」为空，缺少表头行")
            col_map = _match_columns(group, header)
            header_len = len(header)
            rows: list[dict] = []
            for row_idx, row in enumerate(rows_iter, start=2):
                # read_only 模式尾随空单元格会被截断，补齐到表头长度再按列取
                row = tuple(row) + (None,) * (header_len - len(row))
                if all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
                    continue  # 整行空白跳过
                cells = {}
                for idx, key in col_map.items():
                    value = row[idx]
                    cells[key] = value.strip() if isinstance(value, str) else value
                rows.append({"sheet": sheet_name, "row": row_idx, "cells": cells})
            result[group] = rows
        return result
    finally:
        wb.close()


def _match_columns(group: str, header: tuple) -> dict[int, str]:
    """表头行 → {列索引: key}；未知列忽略，日期列缺失或指标列全缺则报错。"""
    col_map: dict[int, str] = {}
    for idx, cell in enumerate(header):
        name = str(cell).strip() if cell is not None else ""
        if name == DATE_HEADER:
            col_map[idx] = "record_date"
        else:
            for key in BODY_GROUPS[group]:
                if name == _column_label(key):
                    col_map[idx] = key
                    break
    if "record_date" not in col_map.values():
        raise ValueError(f"工作表「{SHEET_NAMES[group]}」缺少「{DATE_HEADER}」列")
    if not any(k != "record_date" for k in col_map.values()):
        raise ValueError(f"工作表「{SHEET_NAMES[group]}」未识别到任何指标列")
    return col_map


# ---------- 行级校验 ----------

def validate_rows(parsed: dict[str, list[dict]]) -> tuple[dict[str, list[dict]], list[dict]]:
    """逐行校验 → (有效行 {group: [{record_date, values}]}, 错误明细 [{sheet, row, column, reason}])。

    校验项：日期非空且格式正确、数值可解析、在合理范围内、至少填写一个指标。
    """
    valid: dict[str, list[dict]] = {g: [] for g in IMPORT_GROUPS}
    errors: list[dict] = []
    for group, items in parsed.items():
        for item in items:
            cells = item["cells"]
            row_errors: list[dict] = []

            record_date: date | None = None
            if "record_date" in cells:
                record_date, reason = _parse_date(cells.get("record_date"))
                if reason:
                    row_errors.append(_err(item, DATE_HEADER, reason))

            values: dict[str, float] = {}
            for key in BODY_GROUPS[group]:
                raw = cells.get(key)
                if raw is None:
                    continue
                value, reason = _parse_number(raw)
                if reason:
                    row_errors.append(_err(item, _column_label(key), reason))
                    continue
                lo, hi = RANGES.get(key, (0, 1_000_000))
                if not (lo <= value <= hi):
                    row_errors.append(_err(item, _column_label(key), f"超出合理范围（{lo:g}~{hi:g}）"))
                    continue
                values[key] = value

            if not values:
                row_errors.append(_err(item, "", "至少填写一个指标"))

            if row_errors:
                errors.extend(row_errors)
            else:
                valid[group].append({"record_date": record_date, "values": values})
    return valid, errors


def _err(item: dict, column: str, reason: str) -> dict:
    return {"sheet": item["sheet"], "row": item["row"], "column": column, "reason": reason}


def _parse_date(value) -> tuple[date | None, str | None]:
    """单元格值 → (date, None)；失败返回 (None, 原因)。None/空串按必填缺失处理。"""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, "日期不能为空"
    if isinstance(value, datetime):
        return value.date(), None
    if isinstance(value, date):
        return value, None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # Excel 序列日期号
        try:
            return from_excel(value).date(), None
        except Exception:
            return None, "日期格式应为 YYYY-MM-DD"
    text = str(value).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date(), None
        except ValueError:
            continue
    return None, "日期格式应为 YYYY-MM-DD"


def _parse_number(value) -> tuple[float | None, str | None]:
    """单元格值 → (float, None)；None/空串返回 (None, None) 表示留空；格式错误返回 (None, 原因)。"""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, None
    if isinstance(value, bool):
        return None, "数值格式不正确"
    try:
        return float(value), None
    except (TypeError, ValueError):
        return None, "数值格式不正确"


# ---------- 导出 ----------

def export_workbook(rows_by_group: dict[str, list[dict]]) -> BytesIO:
    """导出 xlsx：两个 Sheet，表头与模板一致，日期升序，可直接再导入。

    rows_by_group[group] = [{"record_date": date, "values": {key: float}}]
    """
    buf = BytesIO()
    wb = Workbook(write_only=True)
    for group in IMPORT_GROUPS:
        ws = wb.create_sheet(SHEET_NAMES[group])
        ws.append(header_of(group))
        rows = sorted(rows_by_group.get(group, []), key=lambda r: r["record_date"])
        for item in rows:
            ws.append([item["record_date"], *[item["values"].get(k) for k in BODY_GROUPS[group]]])
    wb.save(buf)
    return buf
